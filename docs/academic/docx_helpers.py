"""Typography, styling, header/footer, table, and figure primitives for the
TravelAI SRS generator (python-docx)."""
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = '1B365D'
ACCENT = '2E5A8C'
ZEBRA = 'F2F2F2'
CALLOUT_FILL = 'EEF3F8'
BORDER = 'BFBFBF'
MUTED = '555F6A'


def _font_east_asia(style):
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    fonts.set(qn('w:eastAsia'), 'Calibri')


def setup_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    _font_east_asia(normal)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in (('Heading 1', 16, NAVY), ('Heading 2', 13, NAVY),
                              ('Heading 3', 11.5, ACCENT), ('Heading 4', 11, ACCENT)):
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor.from_string(color)
        _font_east_asia(st)
    h1 = doc.styles['Heading 1']
    h1.paragraph_format.page_break_before = True
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(10)
    for nm in ('Heading 2', 'Heading 3', 'Heading 4'):
        doc.styles[nm].paragraph_format.space_before = Pt(10)
        doc.styles[nm].paragraph_format.space_after = Pt(4)


def margins(section):
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)


def field(paragraph, instr, placeholder='1'):
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    begin.set(qn('w:dirty'), 'true')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    run = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = placeholder; run.append(t)
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    r = paragraph.add_run()._r
    for el in (begin, it, sep):
        r.append(el)
    r.append(run)
    r.append(end)


def running_header(section, text):
    section.header.is_linked_to_previous = False
    p = section.header.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.text = ''
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def running_footer(section, left_text):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.text = ''
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left_text + '\t')
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    r2 = p.add_run('Page ')
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor.from_string(MUTED)
    field(p, 'PAGE', '1')
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def cell_borders(cell, color=BORDER, size='4', left=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(left[1]) if (edge == 'left' and left) else size)
        e.set(qn('w:color'), left[0] if (edge == 'left' and left) else color)
        borders.append(e)
    tc_pr.append(borders)


def cell_text(cell, text, *, bold=False, color=None, size=10, align=None, italic=False):
    cell.text = ''
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    for i, line in enumerate(str(text).split('\n')):
        run = p.add_run(('\n' if i else '') + line)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = 'Calibri'
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return cell


def table(doc, headers, rows, widths=None, size=10, header_fill=NAVY, zebra=True):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.allow_autofit = False
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        cell_text(c, h, bold=True, color='FFFFFF', size=size, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(c, header_fill)
        cell_borders(c, color=header_fill, size='4')
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.rows[i + 1].cells[j]
            bold = isinstance(val, str) and val.startswith('**')
            text = val[2:] if bold else val
            cell_text(c, text, bold=bold, size=size,
                      align=WD_ALIGN_PARAGRAPH.CENTER if j == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT)
            if zebra and i % 2 == 1:
                shade(c, ZEBRA)
            cell_borders(c)
    if widths:
        for j, w in enumerate(widths):
            for r in t.rows:
                r.cells[j].width = Inches(w)
    return t


def kv_table(doc, pairs, label_width=1.75, value_width=4.75, size=10):
    t = doc.add_table(rows=len(pairs), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, (k, v) in enumerate(pairs):
        c0, c1 = t.rows[i].cells
        cell_text(c0, k, bold=True, color='FFFFFF', size=size)
        shade(c0, NAVY)
        cell_borders(c0, color=NAVY)
        cell_text(c1, v, size=size)
        shade(c1, ZEBRA if i % 2 else 'FFFFFF')
        cell_borders(c1)
        c0.width, c1.width = Inches(label_width), Inches(value_width)
    return t


def callout(doc, title, text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    c.text = ''
    p = c.paragraphs[0]
    r = p.add_run(title + '  ')
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    shade(c, CALLOUT_FILL)
    cell_borders(c, color='9AB3CC', size='4', left=(NAVY, '24'))
    doc.add_paragraph()


def body(doc, text, size=11, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, color=None, bold=False):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.italic = italic
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p


def bullets(doc, items, size=11):
    for it in items:
        p = doc.add_paragraph(style='List Bullet')
        if isinstance(it, tuple):
            r = p.add_run(it[0])
            r.font.bold = True
            r.font.size = Pt(size)
            r2 = p.add_run(it[1])
            r2.font.size = Pt(size)
        else:
            r = p.add_run(it)
            r.font.size = Pt(size)


def add_figure(doc, path, caption, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(width))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label, text = caption.split(':', 1)
    r1 = cp.add_run(label + ': ')
    r1.font.bold = True
    r1.font.size = Pt(10)
    r2 = cp.add_run(text.strip())
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor.from_string(MUTED)


def toc_field(doc):
    p = doc.add_paragraph()
    field(p, 'TOC \\o "1-3" \\h \\z \\u',
          'Right-click and choose Update Field (or press Ctrl+A then F9) to populate the automatic table of contents.')
